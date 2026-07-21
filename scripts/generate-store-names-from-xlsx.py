#!/usr/bin/env python3
"""One-shot helper: build store_names.json, update question texts, write migration 049."""
from __future__ import annotations

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl required: pip install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path(r"C:\Users\tgaut\Downloads\FM_Store_Name_Number (1).xlsx")
if len(sys.argv) > 1:
    XLSX = Path(sys.argv[1])

wb = openpyxl.load_workbook(XLSX)
rows = list(wb.active.iter_rows(values_only=True))[1:]
names = {str(int(n)): str(name).strip() for name, n in rows if n is not None}

store_names_path = ROOT / "seed" / "store_names.json"
store_names_path.write_text(
    json.dumps(
        {
            "source": XLSX.name,
            "names": names,
        },
        indent=1,
    )
    + "\n",
    encoding="utf-8",
)
print(f"wrote {store_names_path} ({len(names)} stores)")

qs_path = ROOT / "seed" / "question_set_v2.json"
qs = json.loads(qs_path.read_text(encoding="utf-8"))

replacements = {
    "Q5": "Does {{storeName}} have a KOMPASS cart?",
    "Q5a": "Does {{storeName}} put new items on it?",
    "Q5b": "Does {{storeName}} have an area for storing KOMPASS supplies, data skins, pushers, and reset materials?",
    "Q6": "Where at {{storeName}} is the cart or designated area located (if there is one)?",
    "Q7": "Does {{storeName}} have a KOMPASS Champion or reset captain?",
    "Q8": "Does {{storeName}} have cleaning supplies available to use?",
    "Q9": "Does {{storeName}} provide garbage/recycling bags?",
    "Q10": "Does {{storeName}} supply easy access to step ladders and banana boxes?",
    "Q11": "Is our team provided with adequate data skins/channel strips at {{storeName}} to replace broken or dirty ones?",
    "Q12": "Does {{storeName}} have a designated staging area for new shelf-stable grocery items?",
    "Q13": "Does {{storeName}} have a designated staging area for refrigerated/frozen product?",
    "Q14": "Does {{storeName}} have a designated staging area for natural foods items?",
    "Q15": "Does {{storeName}} have a designated staging area for HABA?",
    "Q16": "Does {{storeName}} have a designated staging area for produce?",
    "Q17": "Is our team notified if the Vestcom box at {{storeName}} has been delayed or otherwise unavailable?",
    "Q18": "Is our team informed if there are materials missing from the Vestcom box at {{storeName}}?",
    "Q19": "Are our associates notified if {{storeName}} has not received all shelf strips and POGs for an upcoming reset?",
    "Q20": "Where is the Vestcom box typically found at {{storeName}}?",
    "Q21": "Where are the price tags typically found at {{storeName}}?",
    "Q22": "Is there a regular issue with finding either the box or the tags at {{storeName}}?",
    "Q23": "Does {{storeName}} separate strips and/or tags prior to us performing the resets?",
    "Q24": "Does {{storeName}} have a specific way they want overstock or not-in-set items handled?",
    "Q25": "Does {{storeName}} ask our team to place backstock on the top stock rack?",
    "Q26": "Does {{storeName}} ask us to rotate product and check for out-of-date items?",
    "Q27": "Does {{storeName}} ask us to perform the marking down of items pulled from sets?",
    "Q28": "Does {{storeName}} ask us to make missing tags?",
    "Q29": "Does {{storeName}} ask us to clean fixtures in excess of what would be considered our normal scope of involvement?",
    "Q30": "How do you enter {{storeName}} each morning?",
    "Q31a": "How does the rest of the team enter {{storeName}}?",
    "Q32": "Does anyone at {{storeName}} make it difficult to gain entry?",
    "Q34": "Do you know the name of the store director at {{storeName}}?",
    "Q35": "Do you know the name of the price changer at {{storeName}}?",
    "Q36": "Do you know the name of the grocery manager at {{storeName}}?",
    "Q37": "Do you know the name of the grocery receiver at {{storeName}}?",
    "Q38": "Do any of the following departments at {{storeName}} regularly push back on us performing resets?",
    "Q39": "Does {{storeName}} use DSL (Digital Shelf Labels)?",
    "Q39a": "Does our team get pushback from {{storeName}} regarding use of a handheld (Zebra)?",
    "Q40": "Is the wifi at {{storeName}} acceptable for performing your job?",
    "Q41": "When getting signed out at the end of the day, does someone from {{storeName}} walk the sets to confirm completion standards prior to signing us out?",
    "Q42": "Are there any safety issues or consistent issues with anything or anyone at {{storeName}}?",
    "Q43": "Are there any comments or concerns you would like us to be aware of at {{storeName}}?",
}


def walk(questions):
    for q in questions:
        if q.get("id") in replacements:
            q["text"] = replacements[q["id"]]
        for b in q.get("branches") or []:
            if b.get("id") in replacements:
                b["text"] = replacements[b["id"]]


for sec in qs["sections"]:
    walk(sec.get("questions") or [])

found = set()


def collect(questions):
    for q in questions:
        found.add(q.get("id"))
        for b in q.get("branches") or []:
            found.add(b.get("id"))


for sec in qs["sections"]:
    collect(sec.get("questions") or [])
missing = [k for k in replacements if k not in found]
if missing:
    sys.exit(f"missing question ids: {missing}")

qs_path.write_text(json.dumps(qs, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
print(f"updated {qs_path}")

sql_lines = [
    "-- Survey store names (Division 701 FM name/number list) + store-aware question text",
    "ALTER TABLE survey_store_districts ADD COLUMN IF NOT EXISTS store_name TEXT;",
    "",
]
for num, name in sorted(names.items(), key=lambda x: int(x[0])):
    esc = name.replace("'", "''")
    sql_lines.append(
        f"UPDATE survey_store_districts SET store_name = '{esc}' WHERE store_num = {num};"
    )

spec_json = json.dumps(qs, ensure_ascii=False)
tag = "qspec"
while f"${tag}$" in spec_json:
    tag += "x"
title_esc = qs["title"].replace("'", "''")
sql_lines += [
    "",
    "-- Refresh active question set text with {{storeName}} placeholders",
    (
        f"UPDATE survey_question_sets SET title = '{title_esc}', "
        f"spec = ${tag}${spec_json}${tag}$::jsonb "
        f"WHERE version = {qs['version']};"
    ),
    "",
]

mig = ROOT / "src" / "migrations" / "049_survey_store_names.sql"
mig.write_text("\n".join(sql_lines) + "\n", encoding="utf-8")
print(f"wrote {mig} ({mig.stat().st_size} bytes)")
